





from fastapi import FastAPI, HTTPException, Path, Query, APIRouter
from docx import Document
from .sharepoint_client import SharePointClient
from .ai_client import AIClient
from .doc_processor import (
    extract_placeholders_from_doc,
    extract_table_fields,
    extract_checkboxes,
    generate_filled_document
)
from .schemas import BuildQuestionsRequest, FillRequest
from typing import List, Optional
import uvicorn
import random

app = FastAPI(title="HR Document Auto-Fill API")

# Create clients at startup
sp_client = SharePointClient()
ai_client = AIClient()
router = APIRouter()


# 1. Get all Forms - ONLY RESPONSE CHANGED
@app.get("/forms", response_model=dict)
def list_forms():
    try:
        templates = sp_client.list_docx_templates()
        # Transform templates to HR-style format
        hr_docs = [t.replace(".docx", "") for t in templates]
        return {"unique_HrDocuments": hr_docs}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# 2. Get matching Forms - NEW ENDPOINT (same logic as above with filter)
@app.get("/forms/search", response_model=dict)
def search_forms(query: Optional[str] = Query(None, description="Search query to filter forms")):
    try:
        templates = sp_client.list_docx_templates()
        hr_docs = [t.replace(".docx", "") for t in templates]
        
        # Filter if query provided
        if query:
            query_lower = query.lower()
            hr_docs = [doc for doc in hr_docs if query_lower in doc.lower()]
        
        return {"unique_HrDocuments": hr_docs}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# 3. Build questions - LOGIC UNCHANGED, ONLY RESPONSE FORMAT CHANGED
@app.post("/forms/{template_name}/questions", response_model=List[dict])
def build_questions(template_name: str = Path(..., description="Template file name (including .docx)"),
                    body: BuildQuestionsRequest = None):
    try:
        # Download and inspect doc (UNCHANGED)
        stream = sp_client.download_file(template_name)
        doc = Document(stream)
        placeholders = extract_placeholders_from_doc(doc)
        schedule_fields = extract_table_fields(doc)
        checkbox_options = extract_checkboxes(doc)

        # If client already provided placeholders, use them (UNCHANGED)
        if body and body.placeholders:
            placeholders = body.placeholders

        # Build questions using AI (UNCHANGED)
        qa_pairs = ai_client.build_questions_for_placeholders(placeholders)
        
        # Transform to HR API response format (ONLY THIS CHANGED)
        questions = []
        for p, q in qa_pairs:
            # Handle checkbox_options - it might be a dict or list
            options = []
            if isinstance(checkbox_options, dict):
                options = checkbox_options.get(p, [])
            elif isinstance(checkbox_options, list):
                options = checkbox_options if checkbox_options else []
            
            question_obj = {
                "question": q,
                "options": options,
                "type": "Dropdown" if options else "Single Line Text"
            }
            questions.append(question_obj)

        # Include schedule fields as separate questions
        # Handle schedule_fields - it might be a dict or list
        if isinstance(schedule_fields, dict):
            for field, values in schedule_fields.items():
                questions.append({
                    "question": field,
                    "options": values if values else [],
                    "type": "Dropdown" if values else "Single Line Text"
                })
        elif isinstance(schedule_fields, list) and schedule_fields:
            # If it's a list of field names, add them as text questions
            for field in schedule_fields:
                questions.append({
                    "question": field,
                    "options": [],
                    "type": "Single Line Text"
                })

        return questions
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# 4. Fill template - LOGIC UNCHANGED, ONLY RESPONSE FORMAT CHANGED
@app.post("/forms/{template_name}/fill", response_model=dict)
def fill_template(template_name: str, payload: FillRequest):
    try:
        # Generate document (UNCHANGED LOGIC)
        output_path = generate_filled_document(
            template_name,
            answers=payload.answers or {},
            selected_checkboxes=payload.selected_checkboxes or [],
            schedule_values=payload.schedule_values or {},
            sp_client=sp_client
        )
        
        # Return HR-style submission response (ONLY THIS CHANGED)
        return {
            "status": "success",
            "message": "Request submitted successfully",
            "Request_id": f"HR{random.randint(1000, 9999)}"
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/health")
def health():
    return {"status": "ok"}


# For local run
if __name__ == "__main__":
    uvicorn.run("app.main:app", host="0.0.0.0", port=8000, reload=True)